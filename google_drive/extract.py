from io import BytesIO
import os
import re

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def extract_text(data: bytes, mime_type: str, file_name: str) -> str:
    if mime_type.startswith("text/") or mime_type in {
        "application/json",
        "application/xml",
    }:
        text = data.decode("utf-8-sig", errors="replace")
    elif mime_type == "application/pdf":
        text = _extract_pdf(data)
    elif mime_type == DOCX_MIME_TYPE:
        text = _extract_docx(data)
    elif mime_type == XLSX_MIME_TYPE:
        text = _extract_xlsx(data)
    elif mime_type == PPTX_MIME_TYPE:
        text = _extract_pptx(data)
    else:
        raise ValueError(f"No text extractor for {mime_type} ({file_name})")

    text = _normalize_text(text)
    max_characters = _max_extracted_characters()
    if len(text) > max_characters:
        text = text[:max_characters]
    return text


def chunk_text(
    text: str,
    *,
    max_characters: int = 4_000,
    overlap: int = 300,
) -> list[str]:
    text = text.strip()
    if not text:
        return []
    if max_characters <= overlap:
        raise ValueError("Chunk size must be larger than overlap")

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_characters, len(text))
        if end < len(text):
            minimum_break = start + max_characters // 2
            candidates = (
                text.rfind("\n\n", minimum_break, end),
                text.rfind("\n", minimum_break, end),
                text.rfind(" ", minimum_break, end),
            )
            split_at = max(candidates)
            if split_at > start:
                end = split_at

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)

    return chunks


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            lines.append("\t".join(values))
    return "\n".join(lines)


def _extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            lines.append(f"Sheet: {worksheet.title}")
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append("\t".join(values).rstrip())
    finally:
        workbook.close()
    return "\n".join(lines)


def _extract_pptx(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    lines: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        lines.append(f"Slide {slide_number}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{4,}", "\n\n\n", text).strip()


def _max_extracted_characters() -> int:
    raw_value = os.getenv("GOOGLE_DRIVE_MAX_EXTRACTED_CHARACTERS", "2000000")
    try:
        value = int(raw_value)
    except ValueError as exc:
        raise ValueError(
            "GOOGLE_DRIVE_MAX_EXTRACTED_CHARACTERS must be a whole number"
        ) from exc
    if value <= 0:
        raise ValueError("GOOGLE_DRIVE_MAX_EXTRACTED_CHARACTERS must be greater than zero")
    return value
